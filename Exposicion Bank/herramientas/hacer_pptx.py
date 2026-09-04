"""Arma exposicion_bank.pptx con las figuras de figuras/. Se corre desde la carpeta Exposicion Bank."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x55, 0x55)
FIG = "figuras"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blanco = prs.slide_layouts[6]


def titulo(slide, texto, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = AZUL; p.font.name = "Calibri"
    if sub:
        p2 = tb.text_frame.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(16); p2.font.color.rgb = GRIS; p2.font.name = "Calibri"
    ln = slide.shapes.add_shape(1, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.04))
    ln.fill.solid(); ln.fill.fore_color.rgb = AZUL; ln.line.fill.background()


def vinetas(slide, items, left, top, w, h, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        nivel = 0
        if it.startswith("  "):
            nivel = 1; it = it.strip()
        p.text = ("• " if nivel == 0 else "– ") + it
        p.level = nivel
        p.font.size = Pt(size - 2 * nivel); p.font.name = "Calibri"
        p.space_after = Pt(6)


def imagen(slide, nombre, left, top, w):
    slide.shapes.add_picture(os.path.join(FIG, nombre), Inches(left), Inches(top), width=Inches(w))


def tabla(slide, cab, filas, left, top, w, h, size=13, anchos=None):
    t = slide.shapes.add_table(len(filas) + 1, len(cab), Inches(left), Inches(top), Inches(w), Inches(h)).table
    for j, a in enumerate(anchos or []):
        t.columns[j].width = Inches(a)
    for j, c in enumerate(cab):
        cel = t.cell(0, j); cel.text = c
        for p in cel.text_frame.paragraphs:
            p.font.size = Pt(size); p.font.bold = True; p.font.name = "Calibri"
    for i, fila in enumerate(filas):
        for j, c in enumerate(fila):
            cel = t.cell(i + 1, j); cel.text = str(c)
            for p in cel.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.name = "Calibri"


def pie(slide, texto):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.5))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]; p.text = texto
    p.font.size = Pt(12); p.font.color.rgb = GRIS; p.font.name = "Calibri"


def nueva(texto, sub=None):
    s = prs.slides.add_slide(blanco)
    titulo(s, texto, sub)
    return s


def portada():
    s = prs.slides.add_slide(blanco)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11.3), Inches(3))
    tf = tb.text_frame; tf.word_wrap = True
    lineas = [("Bank Marketing", 44, AZUL, True),
              ("Preparación del dataset paso a paso para la herramienta de clasificación", 24, GRIS, False),
              ("", 12, GRIS, False),
              ("Sergio Duarte", 20, None, False),
              ("Inteligencia Computacional. Maestría en Ciencias de la Información y las Comunicaciones", 16, GRIS, False),
              ("Universidad Distrital Francisco José de Caldas", 16, GRIS, False)]
    for k, (t, sz, col, neg) in enumerate(lineas):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text = t; p.font.size = Pt(sz); p.font.bold = neg; p.font.name = "Calibri"
        if col: p.font.color.rgb = col


portada()

s = nueva("El dataset", "UCI Machine Learning Repository, Bank Marketing (Moro, Cortez y Rita, 2014)")
vinetas(s, [
    "Campañas telefónicas de un banco portugués entre mayo de 2008 y noviembre de 2010",
    "Objetivo. Predecir si el cliente suscribe un depósito a plazo (variable y)",
    "45 211 registros, 16 atributos de entrada y la clase",
    "Sin valores faltantes explícitos, pero varias categóricas usan la etiqueta unknown",
    "Clase muy desbalanceada, 39 922 no frente a 5 289 sí (11.7 % de positivos)",
    "Archivo bank-full.csv, separado por punto y coma",
    "Plan de trabajo en cuatro pasos sobre el dato",
    "  T0 crudo, tal como viene",
    "  T1 arreglo de columnas, filas, faltantes y atípicos",
    "  T2 codificación, estandarización y normalización",
    "  T3 balanceo de clases, por submuestreo y por sobremuestreo",
], 0.7, 1.6, 12, 5.3)

s = nueva("Variables", "Tres grupos según la documentación original")
tabla(s, ["Grupo", "Variable", "Tipo", "Observación"], [
    ["Cliente", "age", "numérica", "18 a 95 años"],
    ["Cliente", "job", "categórica 12 niveles", "incluye unknown (288), quedan 11 niveles"],
    ["Cliente", "marital", "categórica 3 niveles", "divorced agrupa divorciados y viudos"],
    ["Cliente", "education", "categórica ordinal", "primary, secondary, tertiary, unknown (1 857)"],
    ["Cliente", "default, housing, loan", "binarias yes/no", "crédito en mora, hipoteca, préstamo personal"],
    ["Cliente", "balance", "numérica", "saldo medio anual en euros, cola muy larga"],
    ["Último contacto", "contact", "categórica 3 niveles", "unknown en el 29 % de las filas"],
    ["Último contacto", "day, month", "numérica y categórica", "fecha del último contacto sin año"],
    ["Último contacto", "duration", "numérica", "segundos de la llamada, se conoce después de llamar"],
    ["Campaña", "campaign", "numérica", "contactos en esta campaña"],
    ["Campaña", "pdays", "numérica", "días desde la campaña anterior, -1 si nunca"],
    ["Campaña", "previous", "numérica", "contactos antes de esta campaña"],
    ["Campaña", "poutcome", "categórica 4 niveles", "unknown en el 82 % de las filas"],
], 0.5, 1.5, 12.3, 5.3, size=12)

s = nueva("T0. Distribución de las numéricas", "Histogramas de las siete variables numéricas y de la clase")
imagen(s, "fig01.png", 0.5, 1.45, 8.6)
vinetas(s, [
    "age es la única con una distribución aproximadamente simétrica",
    "balance, campaign y previous presentan colas muy pronunciadas",
    "pdays es casi toda -1 (82 % nunca contactados)",
    "duration se concentra en llamadas cortas",
    "day cubre todo el mes de forma casi uniforme",
    "y desbalanceada, 7.5 no por cada sí",
], 9.3, 1.6, 3.7, 5.5, size=16)

s = nueva("T0. Diagramas de caja", "Cada variable en su propia escala")
imagen(s, "fig02.png", 0.5, 1.45, 8.6)
vinetas(s, [
    "balance con valores hasta 102 127 y mínimo de -8 019",
    "campaign con hasta 63 llamadas a un mismo cliente",
    "previous con un caso de 275 contactos",
    "pdays y previous tienen rango intercuartil cero, la regla de Tukey no sirve ahí",
    "age muestra pocos atípicos, mayores de 70 años",
], 9.3, 1.6, 3.7, 5.5, size=16)

s = nueva("T0. Categóricas y valores unknown", "El dataset no trae faltantes, pero unknown cumple ese papel")
imagen(s, "fig03.png", 0.5, 1.45, 7.6)
tabla(s, ["Variable", "unknown", "Decisión"], [
    ["job", "288 (0.6 %)", "se eliminan las filas"],
    ["education", "1 857 (4.1 %)", "se imputa con la moda del oficio"],
    ["contact", "13 020 (28.8 %)", "se conserva como categoría"],
    ["poutcome", "36 959 (81.7 %)", "coincide con nunca contactado"],
], 8.4, 1.6, 4.6, 2.2, size=13)
vinetas(s, [
    "month muy concentrado en mayo",
    "default casi constante, 1.8 % en mora",
    "Las unknown de poutcome coinciden con las filas de previous igual a cero salvo cinco casos",
], 8.4, 4.2, 4.6, 2.5, size=15)

s = nueva("T1. Arreglo de columnas", "Menos columnas y con significado, con intervención mínima sobre el dato")
tabla(s, ["Columna", "Tratamiento", "Por qué"], [
    ["duration", "se elimina", "solo se conoce al terminar la llamada, es fuga de información hacia la clase"],
    ["day y month", "se unen en dia_anio (1 a 365)", "una sola columna numérica en vez de una numérica y una categórica de 12 niveles"],
    ["education", "ordinal 1, 2, 3", "los niveles tienen orden natural, no hace falta one hot"],
    ["education unknown", "moda dentro de cada job", "el oficio es el mejor indicio del nivel educativo"],
    ["job unknown", "se eliminan 288 filas", "no hay con qué imputar y son el 0.6 %"],
    ["contact unknown", "se conserva", "es cerca del 29 % del dato y tiene comportamiento propio"],
    ["pdays", "dos campos, ver siguiente lámina", "el -1 es un código de ausencia, no una cantidad de días"],
], 0.5, 1.5, 12.3, 4.2, size=13)
pie(s, "Para dia_anio se usa un año no bisiesto de referencia porque el dataset no trae el año. Resultado 44 923 filas y 15 columnas de entrada.")

s = nueva("T1. pdays y el umbral de contacto previo", "Tasa de sí según los días desde el último contacto")
imagen(s, "fig04.png", 0.5, 1.45, 7.4)
vinetas(s, [
    "Nunca contactado responde sí el 9.2 % de las veces (línea roja)",
    "Hasta 239 días la tasa se mantiene muy por encima de la base",
    "Entre 240 y 364 días cae a 12 % y 8.6 %, ya se comporta como nunca contactado",
    "Desde 365 días la tasa vuelve a subir a 27 %, pero son 691 filas (1.5 %) y se asumen bajo el mismo umbral para no abrir un tercer caso",
    "Umbral elegido 240 días, unos ocho meses",
    "contactado_antes vale 1 si pdays está entre 0 y 240",
    "pdays se deja en 240 para el resto, así el -1 desaparece",
], 8.1, 1.6, 4.9, 5.5, size=15)

s = nueva("T1. Atípicos", "Regla de 1.5 veces el rango intercuartil, recortando al borde")
imagen(s, "fig05.png", 0.5, 1.45, 7.4)
tabla(s, ["Variable", "Fuera del rango", "Tratamiento"], [
    ["age", "480", "recorte al borde superior 70.5"],
    ["balance", "4 712", "recorte a -1 962 y 3 462"],
    ["campaign", "3 031", "recorte a 6"],
    ["previous", "IQR cero", "recorte al percentil 99 (9)"],
    ["pdays", "IQR cero", "ya acotada por el umbral"],
    ["education, dia_anio", "sin atípicos", "sin tratamiento"],
], 8.1, 1.6, 4.9, 2.7, size=12)
vinetas(s, [
    "Se recorta al borde en vez de borrar filas para no perder varios miles de filas",
    "Los 5 255 registros de la clase sí son escasos y eliminar filas los reduciría todavía más",
], 8.1, 5.3, 4.9, 1.8, size=14)

s = nueva("T2. Codificación", "Todo queda numérico para la herramienta de clasificación")
tabla(s, ["Tipo", "Variables", "Método", "Columnas"], [
    ["binaria", "default, housing, loan, contactado_antes", "0 y 1", "4"],
    ["ordinal", "education", "1 primary, 2 secondary, 3 tertiary", "1"],
    ["nominal", "job", "one hot", "11"],
    ["nominal", "marital", "one hot", "3"],
    ["nominal", "contact", "one hot", "3"],
    ["nominal", "poutcome", "one hot sin unknown", "3"],
    ["numérica", "age, balance, campaign, pdays, previous, dia_anio", "sin cambio", "6"],
], 0.5, 1.5, 12.3, 3.6, size=13)
vinetas(s, [
    "poutcome_unknown se elimina por redundancia, se deduce de las otras tres indicadoras y marca a los nunca contactados, que ya identifica contactado_antes",
    "De 15 columnas de entrada se pasa a 31, todas numéricas",
], 0.7, 5.4, 12, 1.4, size=16)

s = nueva("T2. Estandarización y normalización", "Sobre age, balance, campaign, pdays, previous y dia_anio. Las indicadoras, contactado_antes y education se quedan como están")
imagen(s, "fig06.png", 0.5, 1.45, 8.2)
vinetas(s, [
    "T2z con z score, media cero y desviación uno",
    "T2n con escalado al rango 0 a 1",
    "pdays se ve concentrada porque el 90 % de las filas vale 240 y el resto queda por debajo",
    "previous sigue concentrada en cero, es la variable con menos información",
    "Se guardan las dos versiones para comparar en la herramienta",
], 8.9, 1.6, 4.1, 5.5, size=15)

s = nueva("T3. Balanceo por dos caminos", "Submuestreo de la clase no (T3_1) y sobremuestreo sintético de la clase sí (T3_2)")
imagen(s, "fig07.png", 0.5, 1.45, 7.4)
vinetas(s, [
    "T3_1. Se toman al azar 5 255 no para igualar a los 5 255 sí, quedan 10 510 filas",
    "T3_2. Se crean 34 413 sí sintéticos entre cada positivo y uno de sus cinco vecinos más cercanos, quedan 79 336 filas",
    "En T3_2 las indicadoras se redondean para que sigan valiendo 0 o 1 y education para que siga en 1, 2 o 3",
    "No se duplican filas porque la copia y su original caen en particiones distintas y la prueba deja de ser independiente",
    "Con los sintéticos el riesgo baja pero no desaparece, la comparación entre T3_1 y T3_2 debe leerse con esa reserva",
    "Semilla fija en los dos caminos y las mismas filas en la versión z y en la versión n",
], 8.1, 1.6, 4.9, 5.5, size=15)

s = nueva("Categóricas y numéricas mezcladas", "Efecto sobre la distancia en KNN y sobre el sobremuestreo")
vinetas(s, [
    "Tras T2 todo es numérico, las one hot y las binarias valen 0 o 1, education va de 1 a 3 y las seis numéricas están en z score o en rango 0 a 1",
    "Dos clientes con distinto oficio difieren en dos columnas one hot, así que la distancia euclídea suma un salto fijo por cada categórica distinta",
    "Con z score las numéricas pesan más que las indicadoras, con rango 0 a 1 quedan en la misma escala, por eso se comparan las versiones z y n",
    "KNN sirve, pero conviene la distancia euclídea sobre la versión n, o la distancia de Manhattan, que trata cada columna por igual",
    "Los árboles y los modelos de conjunto no dependen de la escala y manejan las indicadoras sin problema",
    "En T3_2 se redondean las indicadoras y education para no crear clientes con valores intermedios en oficio, estado civil o nivel educativo, que no tienen interpretación",
], 0.7, 1.6, 12, 5.5, size=17)

s = nueva("Resumen de las versiones del dato", "Tablas que quedan en bank_pasos.mat")
tabla(s, ["Paso", "Filas", "Columnas de entrada", "No", "Sí", "Qué cambió"], [
    ["T0", "45 211", "16", "39 922", "5 289", "dato crudo"],
    ["T1", "44 923", "15", "39 668", "5 255", "sin duration, dia_anio, education ordinal, pdays con umbral, atípicos recortados"],
    ["T2z y T2n", "44 923", "31", "39 668", "5 255", "one hot, binarias en 0 y 1, z score o rango 0 a 1"],
    ["T3_1z y T3_1n", "10 510", "31", "5 255", "5 255", "submuestreo de la clase no"],
    ["T3_2z y T3_2n", "79 336", "31", "39 668", "39 668", "sobremuestreo sintético de la clase sí"],
], 0.5, 1.5, 12.3, 3.0, size=14, anchos=[1.6, 1.2, 2.0, 1.2, 1.2, 5.1])
vinetas(s, [
    "Un solo live script, procesar_bank.mlx, genera las cinco versiones y las figuras",
    "Cada tabla lleva la clase y como última columna",
    "Ninguna versión modifica la clase ni elimina positivos reales, salvo las 34 filas con job unknown",
], 0.7, 5.2, 12, 1.8, size=16)

s = nueva("Siguiente paso. Herramienta de clasificación", "Qué se va a comparar en Classification Learner")
vinetas(s, [
    "Importar cada tabla desde el workspace con y como respuesta",
    "Validación cruzada de 5 particiones en todas las pruebas",
    "El escalado y el sobremuestreo se hicieron sobre todo el conjunto porque la partición se delega a la herramienta, así que las cifras de T3_2 salen algo optimistas",
    "Comparar T2z, T3_1z y T3_2z para ver el efecto de cada balanceo sobre la sensibilidad de la clase sí",
    "Comparar T2z frente a T2n para evaluar la sensibilidad de cada modelo al escalado",
    "Mirar la exactitud, pero sobre todo la matriz de confusión, porque con 88 % de la clase no la exactitud resulta poco informativa",
    "Referencias",
    "  Moro, S., Cortez, P. y Rita, P. (2014). A data driven approach to predict the success of bank telemarketing. Decision Support Systems, 62, 22 a 31",
    "  Tukey, J. W. (1977). Exploratory Data Analysis. Addison Wesley",
], 0.7, 1.6, 12, 5.5, size=17)

prs.save("exposicion_bank.pptx")
print("Listo exposicion_bank.pptx,", len(prs.slides), "laminas")
