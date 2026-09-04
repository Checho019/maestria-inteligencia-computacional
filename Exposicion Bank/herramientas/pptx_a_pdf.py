"""Dibuja exposicion_bank.pptx en un PDF, una pagina por lamina. Se corre desde la carpeta Exposicion Bank."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

pdfmetrics.registerFont(TTFont("Calibri", r"C:\Windows\Fonts\calibri.ttf"))
pdfmetrics.registerFont(TTFont("Calibri-Bold", r"C:\Windows\Fonts\calibrib.ttf"))

PT = 1 / 12700.0
AZUL = (0x1F / 255, 0x3A / 255, 0x5F / 255)
NEGRO = (0, 0, 0)


def fuente(negrita):
    return "Calibri-Bold" if negrita else "Calibri"


def color_de(font, defecto=NEGRO):
    try:
        if font.color and font.color.rgb is not None:
            rgb = font.color.rgb
            return (rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)
    except AttributeError:
        pass
    return defecto


def parrafos(text_frame, tam_defecto=18):
    """Devuelve (texto, tamano, negrita, color, sangria) por parrafo."""
    salida = []
    for p in text_frame.paragraphs:
        texto = "".join(r.text for r in p.runs) or p.text
        runs = p.runs
        f = runs[0].font if runs else p.font
        tam = (f.size or p.font.size)
        tam = tam.pt if tam else tam_defecto
        negrita = bool(f.bold or p.font.bold)
        salida.append((texto, tam, negrita, color_de(f), (p.level or 0) * 18))
    return salida


def dibujar_texto(c, x, y_top, w, lineas_p, alto_max=None):
    """Escribe parrafos desde y_top hacia abajo. Devuelve la altura usada."""
    y = y_top
    for texto, tam, negrita, col, sangria in lineas_p:
        lead = tam * 1.2
        if not texto:
            y -= lead
            continue
        c.setFont(fuente(negrita), tam)
        c.setFillColorRGB(*col)
        for ln in simpleSplit(texto, fuente(negrita), tam, w - sangria - 4):
            y -= lead
            c.drawString(x + 2 + sangria, y, ln)
        y -= 4
    return y_top - y


def dibujar_tabla(c, sh, alto_pagina):
    t = sh.table
    x0 = sh.left * PT
    y = alto_pagina - sh.top * PT
    anchos = [col.width * PT for col in t.columns]
    for i, fila in enumerate(t.rows):
        celdas = []
        alto = 0
        for j, cel in enumerate(fila.cells):
            lp = parrafos(cel.text_frame, 13)
            lineas = []
            for texto, tam, negrita, col, _ in lp:
                for ln in simpleSplit(texto, fuente(negrita), tam, anchos[j] - 8):
                    lineas.append((ln, tam, negrita))
            h = sum(tm * 1.2 for _, tm, _ in lineas) + 8
            alto = max(alto, h)
            celdas.append(lineas)
        x = x0
        for j, lineas in enumerate(celdas):
            if i == 0:
                c.setFillColorRGB(*AZUL)
                c.rect(x, y - alto, anchos[j], alto, stroke=0, fill=1)
            elif i % 2 == 0:
                c.setFillColorRGB(0.93, 0.95, 0.98)
                c.rect(x, y - alto, anchos[j], alto, stroke=0, fill=1)
            c.setStrokeColorRGB(0.75, 0.75, 0.75)
            c.rect(x, y - alto, anchos[j], alto, stroke=1, fill=0)
            yy = y - 4
            for ln, tam, negrita in lineas:
                yy -= tam * 1.2
                c.setFont(fuente(negrita), tam)
                c.setFillColorRGB(*((1, 1, 1) if i == 0 else NEGRO))
                c.drawString(x + 4, yy + tam * 0.25, ln)
            x += anchos[j]
        y -= alto


def convertir(entrada, salida):
    prs = Presentation(entrada)
    W, H = prs.slide_width * PT, prs.slide_height * PT
    c = canvas.Canvas(salida, pagesize=(W, H))
    c.setTitle("Bank Marketing. Preparación del dataset paso a paso")
    c.setAuthor("Sergio Duarte")
    for s in prs.slides:
        for sh in s.shapes:
            x, y_top, w = sh.left * PT, H - sh.top * PT, sh.width * PT
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = ImageReader(io.BytesIO(sh.image.blob))
                c.drawImage(img, x, y_top - sh.height * PT, w, sh.height * PT)
            elif sh.has_table:
                dibujar_tabla(c, sh, H)
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not sh.has_text_frame:
                c.setFillColorRGB(*AZUL)
                c.rect(x, y_top - sh.height * PT, w, sh.height * PT, stroke=0, fill=1)
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not sh.text_frame.text.strip():
                c.setFillColorRGB(*AZUL)
                c.rect(x, y_top - sh.height * PT, w, sh.height * PT, stroke=0, fill=1)
            elif sh.has_text_frame:
                dibujar_texto(c, x, y_top - 4, w, parrafos(sh.text_frame))
        c.showPage()
    c.save()
    print("Listo", salida, "con", len(prs.slides), "paginas")


if __name__ == "__main__":
    convertir("exposicion_bank.pptx", "exposicion_bank.pdf")
