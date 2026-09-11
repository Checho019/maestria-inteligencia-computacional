"""Agrega al final de exposicion_bank_vs2.pptx la lamina de resultados del balanceo y la de conclusiones.
No modifica ninguna lamina existente. Uso  python herramientas/agregar_laminas_vs2.py entrada.pptx salida.pptx"""
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
GRIS = RGBColor(0x55, 0x55, 0x55)
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

entrada, salida = sys.argv[1], sys.argv[2]
prs = Presentation(entrada)
blanco = next(l for l in prs.slide_layouts if l.name == "Blank")


def nueva(texto, sub):
    s = prs.slides.add_slide(blanco)
    tb = s.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = texto
    p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = AZUL; p.font.name = "Calibri"
    p2 = tb.text_frame.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(16); p2.font.color.rgb = GRIS; p2.font.name = "Calibri"
    ln = s.shapes.add_shape(1, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.04))
    ln.fill.solid(); ln.fill.fore_color.rgb = AZUL; ln.line.fill.background()
    return s


def vinetas(s, items, left, top, w, h, size):
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(size); p.font.name = "Calibri"
        p.space_after = Pt(8)


def tabla(s, cab, filas, left, top, w, h, size, anchos):
    t = s.shapes.add_table(len(filas) + 1, len(cab), Inches(left), Inches(top), Inches(w), Inches(h)).table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate([cab] + filas):
        for j, c in enumerate(fila):
            cel = t.cell(i, j); cel.text = c
            for p in cel.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.bold = i == 0; p.font.name = "Calibri"


s = nueva("Resultados del balanceo", "Matriz de confusión de T3A_2 en validación cruzada y comparación con el dato sin balancear")
s.shapes.add_picture(os.path.join(BASE, "figuras", "fig08_matriz_t3_2.png"), Inches(0.5), Inches(1.5), height=Inches(5.5))
tabla(s, ["Versión", "Modelo", "Exactitud", "Sensibilidad", "Especificidad"], [
    ["T2A sin balanceo", "Boosted Trees", "89.4 %", "15.6 %", "99.1 %"],
    ["T3A_2 sobremuestreo", "Fine KNN", "90.8 %", "96.9 %", "84.7 %"],
], 6.1, 1.6, 6.8, 1.2, 14, [1.9, 1.4, 1.1, 1.2, 1.2])
vinetas(s, [
    "Con sobremuestreo el modelo acierta 38 439 de 39 668 sí, frente a 818 de 5 255 sin balancear",
    "El costo recae en la clase no, 6 088 clientes que no compran quedan marcados como compradores",
    "Fine KNN usa un solo vecino y cada sí sintético nace entre dos positivos reales, así que en la validación casi siempre encuentra a su origen. La cifra es optimista",
    "Con submuestreo (T3A_1) la sensibilidad sube a cerca del 56 % sin ese sesgo, con 10 510 filas",
], 6.1, 3.2, 6.8, 3.9, 16)

s = nueva("Conclusiones", "Lectura conjunta de los dos tratamientos")
vinetas(s, [
    "La exactitud no sirve como criterio único. Predecir siempre no ya da 88.3 % y casi todos los modelos quedan entre 89 % y 91 %",
    "duration explica buena parte del desempeño. Con ella, en T0 y en el segundo tratamiento, la sensibilidad llega a 49.8 % y 55.2 %. Sin ella cae a 15.6 %. Como solo se conoce al terminar la llamada, no sirve para decidir a quién llamar",
    "La codificación y el escalado no cambian el resultado de los árboles, T1A y T2A dan la misma exactitud y la misma sensibilidad",
    "El balanceo es lo que más recupera la clase sí sin usar duration. El submuestreo alcanza cerca del 56 %, un nivel similar al del segundo tratamiento con duration",
    "El sobremuestreo hecho antes de partir infla la validación. Debe aplicarse solo dentro del entrenamiento de cada partición y medirse sobre un conjunto de prueba reservado",
    "Para una campaña real conviene un modelo sin duration, entrenado con balanceo y evaluado por sensibilidad y precisión",
], 0.7, 1.6, 12, 5.5, 18)

prs.save(salida)
print("Listo", salida, "con", len(prs.slides), "laminas")
